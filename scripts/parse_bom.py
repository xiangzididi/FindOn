"""Parse a small BOM file from stdin and emit JSON to stdout."""

from __future__ import annotations

import csv
import io
import json
import re
import sys
from pathlib import Path

MAX_ITEMS = 300
NAME_HEADERS = {"name", "item", "part", "component", "名称", "物料", "物料名称", "配件", "配件名称", "器件", "器件名称", "零件", "零件名称"}
MODEL_HEADERS = {"model", "spec", "sku", "型号", "规格", "规格型号", "物料编码", "编码"}
QTY_HEADERS = {"qty", "quantity", "count", "amount", "数量", "个数", "件数", "需求数量"}


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def normalized(value: object) -> str:
    return re.sub(r"[\s_\-—()（）\[\]【】:：./]+", "", clean(value).lower())


def quantity(value: object) -> int:
    match = re.search(r"\d+(?:\.\d+)?", clean(value))
    return max(1, min(int(float(match.group())), 9999)) if match else 1


def rows_to_items(rows: list[list[object]]) -> list[dict[str, object]]:
    prepared = [[clean(cell) for cell in row] for row in rows]
    prepared = [row for row in prepared if any(row)]
    if not prepared:
        return []
    header = name_col = model_col = qty_col = None
    for row_index, row in enumerate(prepared[:12]):
        values = [normalized(cell) for cell in row]
        name_col = next((i for i, value in enumerate(values) if value in NAME_HEADERS), None)
        if name_col is not None:
            header = row_index
            model_col = next((i for i, value in enumerate(values) if value in MODEL_HEADERS), None)
            qty_col = next((i for i, value in enumerate(values) if value in QTY_HEADERS), None)
            break
    output: list[dict[str, object]] = []
    start = header + 1 if header is not None else 0
    for row_number, row in enumerate(prepared[start:], start=start + 1):
        if len(output) >= MAX_ITEMS:
            break
        values = [value for value in row if value]
        if not values:
            continue
        if name_col is not None:
            name = row[name_col] if name_col < len(row) else ""
            model = row[model_col] if model_col is not None and model_col < len(row) else ""
            qty = quantity(row[qty_col]) if qty_col is not None and qty_col < len(row) else 1
            if model and normalized(model) != normalized(name):
                name = f"{name} {model}".strip()
        else:
            name = values[1] if len(values) > 1 and values[0].isdigit() else values[0]
            inline = re.match(r"^(.*?)(?:\s*[xX×*]\s*|\s+)(\d+(?:\.\d+)?)\s*(?:个|件|只|条|pcs?)?$", name, re.I)
            if inline:
                name, qty = inline.group(1), quantity(inline.group(2))
            else:
                numbers = [value for value in values[1:] if re.fullmatch(r"\d+(?:\.\d+)?", value)]
                qty = quantity(numbers[-1]) if numbers else 1
        name = clean(name).strip(" -_:：")
        if not name or normalized(name) in NAME_HEADERS:
            continue
        output.append({"name": name[:160], "quantity": qty, "sourceRow": row_number})
    merged: dict[str, dict[str, object]] = {}
    for item in output:
        key = normalized(item["name"])
        if key in merged:
            merged[key]["quantity"] = min(9999, int(merged[key]["quantity"]) + int(item["quantity"]))
        else:
            merged[key] = item
    return list(merged.values())


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError("text_encoding_not_supported")


def parse_file(filename: str, data: bytes) -> list[dict[str, object]]:
    extension = Path(filename).suffix.lower()
    rows: list[list[object]] = []
    if extension == ".xlsx":
        if not data.startswith(b"PK"):
            raise ValueError("invalid_xlsx_file")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    rows.append(list(row))
                    if len(rows) >= MAX_ITEMS + 20:
                        break
        finally:
            workbook.close()
    elif extension == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            for sheet in workbook.sheets():
                for index in range(min(sheet.nrows, MAX_ITEMS + 20 - len(rows))):
                    rows.append(sheet.row_values(index))
        finally:
            workbook.release_resources()
    elif extension in {".csv", ".txt"}:
        text = decode_text(data)
        if extension == ".csv":
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = [list(row) for row in csv.reader(io.StringIO(text), dialect)]
        else:
            rows = [[line] for line in text.splitlines() if line.strip()]
    else:
        raise ValueError("unsupported_file_type")
    items = rows_to_items(rows)
    if not items:
        raise ValueError("no_bom_items_found")
    return items


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    filename = Path(sys.argv[1]).name if len(sys.argv) > 1 else ""
    try:
        data = sys.stdin.buffer.read(6 * 1024 * 1024)
        items = parse_file(filename, data)
        print(json.dumps({"items": items, "count": len(items), "lines": [f"{item['name']} ×{item['quantity']}" for item in items]}, ensure_ascii=False))
        return 0
    except (ImportError, ModuleNotFoundError):
        print(json.dumps({"error": "spreadsheet_parser_unavailable"}))
    except ValueError as error:
        print(json.dumps({"error": str(error)}))
    except Exception:
        print(json.dumps({"error": "file_parse_failed"}))
    return 2


if __name__ == "__main__":
    raise SystemExit(main())

